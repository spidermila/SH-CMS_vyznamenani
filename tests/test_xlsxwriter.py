import os

import openpyxl  # type: ignore
import pandas as pd  # type: ignore
import pytest  # type: ignore
from conftest import cell  # type: ignore
from conftest import YEAR  # type: ignore

import libs.xlsxwriter
from libs.xlsxwriter import Xlsxwriter

GREY = '00969696'
GREEN = '00CCFFCC'


def write_and_load(hasici, path):
    Xlsxwriter(hasici).write(str(path), YEAR)
    return openpyxl.load_workbook(path)['splnene podminky']


def test_write(make_hasic, tmp_path):
    h = make_hasic(cestne_uznani_sdh=2000)
    h.jmeno = 'Karel Zkusebni'
    ws = write_and_load([h], tmp_path / 'o.xlsx')
    assert ws.freeze_panes == 'A2'
    assert ws['A1'].value is None
    assert ws['A2'].value == 'Karel Zkusebni'

    obdrzene = cell(ws, 'Čestné uznání SDH')
    assert obdrzene.value == 'Obdrženo 2000'
    assert obdrzene.fill.start_color.rgb == GREY

    nedava = cell(ws, 'Čestné uznání okrsku')
    assert nedava.value == 'Již se nedává'
    assert nedava.fill.start_color.rgb == GREY

    splneno = cell(ws, 'Čestné uznání OSH - MSH')
    assert splneno.value == 'Podmínky splněny'
    assert splneno.fill.start_color.rgb == GREEN

    prerekvizita = cell(ws, 'Čestné uznání KSH - MSH')
    assert prerekvizita.value == 'Čestné uznání OSH - MSH, Medaile Za zásluhy'
    assert prerekvizita.fill.fill_type is None


def test_permission_error_retry(make_hasic, tmp_path, monkeypatch, capsys):
    real = os.replace
    calls = []

    def fake(*args):
        calls.append(args)
        if len(calls) == 1:
            raise PermissionError
        return real(*args)

    monkeypatch.setattr(libs.xlsxwriter.os, 'replace', fake)
    monkeypatch.setattr('builtins.input', lambda: '')
    h = make_hasic()
    h.jmeno = 'Karel Zkusebni'
    out = tmp_path / 'o.xlsx'
    ws = write_and_load([h], out)
    assert len(calls) == 2
    assert capsys.readouterr().out.count('je používán jinou aplikací. Zavřete') == 1
    assert ws['A2'].value == 'Karel Zkusebni'


@pytest.mark.parametrize('existing', [False, True])
def test_write_failure(make_hasic, tmp_path, monkeypatch, existing):
    def boom(*args, **kwargs):
        raise RuntimeError('boom')

    monkeypatch.setattr(pd.DataFrame, 'to_excel', boom)
    h = make_hasic()
    h.jmeno = 'Karel Zkusebni'
    out = tmp_path / 'o.xlsx'
    if existing:
        out.write_bytes(b'old')
    with pytest.raises(RuntimeError, match='boom'):
        Xlsxwriter([h]).write(str(out), YEAR)
    assert os.listdir(tmp_path) == (['o.xlsx'] if existing else [])
    if existing:
        assert out.read_bytes() == b'old'


def test_cleanup_failure_does_not_mask(make_hasic, tmp_path, monkeypatch):
    def fail(*args):
        raise OSError('replace')

    monkeypatch.setattr(libs.xlsxwriter.os, 'replace', fail)
    h = make_hasic()
    h.jmeno = 'Karel Zkusebni'
    with pytest.raises(OSError, match='replace'):
        Xlsxwriter([h]).write(str(tmp_path / 'o.xlsx'), YEAR)
    assert os.listdir(tmp_path) == []


def test_cleanup_error_does_not_mask(make_hasic, tmp_path, monkeypatch):
    real = os.remove

    def replace(*args):
        raise OSError('replace')

    def remove(path):
        real(path)
        if os.path.dirname(path) == str(tmp_path):
            raise OSError('remove')

    monkeypatch.setattr(libs.xlsxwriter.os, 'replace', replace)
    monkeypatch.setattr(libs.xlsxwriter.os, 'remove', remove)
    h = make_hasic()
    h.jmeno = 'Karel Zkusebni'
    with pytest.raises(OSError, match='replace'):
        Xlsxwriter([h]).write(str(tmp_path / 'o.xlsx'), YEAR)
