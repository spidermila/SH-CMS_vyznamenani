import io

import openpyxl  # type: ignore
import pytest  # type: ignore

from libs.processor import PrazdnyExportError
from libs.processor import process_txt_content

EXPORT = '''Jan Novák
==========================
- 01.02.1970, Datum narození
- 55, Věk
- 01.01.1990, Členem SHČMS od
- 2005, Čestné uznání SDH
konec
'''


def test_process_txt_content():
    ws = openpyxl.load_workbook(
        io.BytesIO(process_txt_content(EXPORT)),
    )['splnene podminky']
    assert ws['A2'].value == 'Jan Novák'
    assert ws.max_row == 2


@pytest.mark.parametrize(
    'txt', ['', 'konec\n', 'Jan\n' + '=' * 26 + '\n- 55, Věk\nkonec\n'],
)
def test_empty_export(txt):
    with pytest.raises(PrazdnyExportError):
        process_txt_content(txt)
