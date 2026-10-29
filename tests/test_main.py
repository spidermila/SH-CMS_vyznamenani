import os
import runpy
import sys
from pathlib import Path

import openpyxl  # type: ignore
import pytest  # type: ignore
from conftest import cell  # type: ignore

SCRIPT = str(Path(__file__).parents[1] / 'SH-CMS_vyznamenani.py')


def run(monkeypatch, *args):
    monkeypatch.setattr(sys, 'argv', ['x', *args])
    return runpy.run_path(SCRIPT)['main']()


def test_missing_file(monkeypatch, capsys, tmp_path):
    missing = str(tmp_path / 'nope.txt')
    assert run(monkeypatch, missing, '-o', str(tmp_path / 'o.xlsx')) == 1
    assert f'Soubor {missing} neexistuje!' in capsys.readouterr().out


@pytest.mark.parametrize('name', ['out.xls', 'out.txt'])
def test_bad_extension(monkeypatch, capsys, export_file, tmp_path, name):
    out = tmp_path / name
    assert run(monkeypatch, str(export_file), '-o', str(out)) == 1
    assert 'Výstupní soubor musí mít příponu .xlsx!' in (
        capsys.readouterr().out
    )
    assert not out.exists()


@pytest.mark.parametrize('name', ['OUT.XLSX', 'out.XlSx'])
def test_uppercase_extension(monkeypatch, export_file, tmp_path, name):
    folder = tmp_path / 'a'
    folder.mkdir()
    assert run(monkeypatch, str(export_file), '-o', str(folder / name)) == 0
    assert 'splnene podminky' in openpyxl.load_workbook(folder / name)
    assert os.listdir(folder) == [name]


@pytest.mark.parametrize('rok', [['-r', '2025'], []])
def test_happy_path(monkeypatch, export_file, tmp_path, rok):
    out = tmp_path / 'o.xlsx'
    args = [str(export_file), '-o', str(out), *rok]
    assert run(monkeypatch, *args) == 0
    ws = openpyxl.load_workbook(out)['splnene podminky']
    assert ws['A2'].value == 'Karel Zkusebni'
    assert ws['A3'].value == 'Jana Vymyslena'
    assert ws.freeze_panes == 'A2'


@pytest.mark.parametrize(
    ('rok', 'text'), [
        ('2025', '2035'), ('2040', 'Podmínky splněny'),
    ],
)
def test_rok_option(monkeypatch, export_file, tmp_path, rok, text):
    out = tmp_path / 'o.xlsx'
    run(monkeypatch, str(export_file), '-o', str(out), '-r', rok)
    ws = openpyxl.load_workbook(out)['splnene podminky']
    assert cell(ws, 'Stužka za věrnost 20 let', row=3).value == text


def test_empty_export(monkeypatch, capsys, tmp_path):
    empty = tmp_path / 'empty.txt'
    empty.write_text('')
    out = tmp_path / 'o.xlsx'
    assert run(monkeypatch, str(empty), '-o', str(out)) == 1
    assert 'Export neobsahuje' in capsys.readouterr().out
    assert not out.exists()


def test_main_guard(monkeypatch, tmp_path):
    monkeypatch.setattr(sys, 'argv', ['x', str(tmp_path / 'nope.txt')])
    with pytest.raises(SystemExit) as exc:
        runpy.run_path(SCRIPT, run_name='__main__')
    assert exc.value.code == 1
